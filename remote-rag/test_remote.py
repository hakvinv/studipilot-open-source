import base64
import hashlib
import json
import os
import re
import tempfile
import unittest
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlparse

from starlette.testclient import TestClient
from server import create_server
import core
import numpy as np
from openpyxl import Workbook

URL='https://uni.example'
OWNER='synthetic-owner-secret'
SERVICE='synthetic-service-secret'


class RemoteTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)
        self.env=patch.dict(os.environ,OWNER_SECRET_SHA256=hashlib.sha256(OWNER.encode()).hexdigest(),
                            SERVICE_TOKEN_SHA256=hashlib.sha256(SERVICE.encode()).hexdigest())
        self.env.start();self.addCleanup(self.env.stop)
        self.app=create_server(self.path,URL,False)
        self.client=TestClient(self.app,base_url=URL)
        self.client.__enter__();self.addCleanup(self.client.__exit__,None,None,None)
        self.key=hashlib.sha256(b'01_StudIP/Test.pdf').hexdigest()
        cache=self.path/'corpus';(cache/'documents').mkdir(parents=True)
        (cache/'documents'/(self.key+'.json')).write_text(json.dumps({
            'id':self.key,'source':'01_StudIP/Test.pdf','course':'Ökonometrie Fortgeschrittene [12345678]',
            'sha256':'synthetic','chunks':[{'page':2,'text':'Instrumentvariablen loesen das Endogenitaetsproblem.'}]}))
        (cache/'catalogue.json').write_text(json.dumps({'files':{'01_StudIP/Test.pdf':{'id':self.key,'status':'indexed'}},
            'status':{'updated_at':'2026-10-19T00:00:00Z','pending':2,'review_required':1,'complete':False}}))
        core.backend.ingest(cache)

    def rpc(self,token=None,method='tools/list',params=None):
        headers={'Accept':'application/json, text/event-stream','Content-Type':'application/json'}
        if token:headers['Authorization']='Bearer '+token
        return self.client.post('/mcp',headers=headers,json={'jsonrpc':'2.0','id':1,'method':method,'params':params or {}})

    def register(self,redirect='https://client.example/callback'):
        r=self.client.post('/register',json={'client_name':'Synthetic client','redirect_uris':[redirect],
            'grant_types':['authorization_code','refresh_token'],'token_endpoint_auth_method':'client_secret_post','scope':'uni:read'})
        self.assertEqual(r.status_code,201,r.text)
        return r.json()

    def authorization(self,client):
        verifier='v'*64
        challenge=base64.urlsafe_b64encode(hashlib.sha256(verifier.encode()).digest()).rstrip(b'=').decode()
        r=self.client.get('/authorize',params={'client_id':client['client_id'],'redirect_uri':'https://client.example/callback',
            'response_type':'code','code_challenge':challenge,'code_challenge_method':'S256','state':'test-state',
            'scope':'uni:read','resource':URL+'/mcp'},follow_redirects=False)
        self.assertIn(r.status_code,[302,303,307],r.text)
        consent=self.client.get(r.headers['location'])
        self.assertEqual(consent.status_code,200)
        self.assertEqual(consent.headers['referrer-policy'],'same-origin')
        self.assertIn("form-action 'self' https://client.example;",consent.headers['content-security-policy'])
        ticket=re.search('name="id" value="([^"]+)"',consent.text).group(1)
        csrf=re.search('name="csrf" value="([^"]+)"',consent.text).group(1)
        form={'id':ticket,'csrf':csrf,'owner_secret':OWNER,'scopes':['uni:read']}
        return verifier,form

    def exchange(self,client,code,verifier):
        return self.client.post('/token',data={'grant_type':'authorization_code','client_id':client['client_id'],
            'client_secret':client['client_secret'],'code':code,'code_verifier':verifier,
            'redirect_uri':'https://client.example/callback','resource':URL+'/mcp'})

    def test_anonymous_invalid_token_and_query_token_rejected(self):
        self.assertEqual(self.client.get('/health').json(),{'status':'ok'})
        self.assertEqual(self.rpc().status_code,401)
        self.assertEqual(self.rpc('invalid').status_code,401)
        r=self.client.post('/mcp?token='+SERVICE,headers={'Accept':'application/json, text/event-stream'},json={})
        self.assertEqual(r.status_code,401)
        self.assertNotIn('Ökonometrie',self.client.get('/').text)

    def test_remote_tools_search_and_original_page(self):
        r=self.rpc(SERVICE)
        self.assertEqual(r.status_code,200,r.text)
        names={x['name'] for x in r.json()['result']['tools']}
        self.assertEqual(names,{'uni_status','uni_courses','uni_search','uni_documents','uni_read_page',
                                'uni_downloader_status','uni_study_plan','uni_calendar_candidates','search','fetch'})
        r=self.rpc(SERVICE,'tools/call',{'name':'uni_search','arguments':{'question':'Instrumentvariablen','course':'Okonometrie Fortgeschrittene'}})
        result=r.json()['result']['structuredContent']
        self.assertEqual(result['sources'][0]['page'],2)
        self.assertFalse(result['coverage_complete'])
        page=self.app.state.corpus.page(self.key,2)
        self.assertIn('Endogenitaet',page['text'])
        self.assertTrue(page['complete'])
        with self.assertRaises(ValueError):self.app.state.corpus.page(self.key,999)

    def test_search_reports_safe_embedding_api_failure_details(self):
        corpus=self.app.state.corpus
        failure=core.embedding_api.EmbeddingError('Embedding-API: HTTP 401. private-body-must-not-leak',
                                                    http_status=401,api_code='invalid_api_key')
        with patch.object(corpus,'status',return_value={'retrieval':'hybrid','warnings':[]}), \
             patch.object(core.embedding_api,'semantic',side_effect=failure), \
             patch.dict(os.environ,OPENAI_API_KEY='synthetic-key'):
            result=corpus.search('Instrumentvariablen')
        self.assertEqual(result['retrieval'],'fulltext')
        self.assertEqual(result['embedding_query'],{'state':'error','http_status':401,
                                                     'api_code':'invalid_api_key','retry_after':None})
        self.assertNotIn('private-body-must-not-leak',json.dumps(result))

    def test_study_plan_is_read_only_and_reports_formula_cache_state(self):
        workbook=Workbook();sheet=workbook.active;sheet.title='Planung'
        sheet['A1']='Bachelorziel';sheet['B1']=2.8;sheet['A2']='Quote';sheet['B2']='=1+1'
        buffer=BytesIO();workbook.save(buffer);payload=buffer.getvalue()
        with patch.object(core.backend,'cloud',return_value=payload) as cloud:
            plan=self.app.state.corpus.study_plan()
        cloud.assert_called_once_with('cat','uni:Desktop/Uni-Management/Stundenplan_WiSe_2026-27.xlsx',timeout=60)
        self.assertTrue(plan['read_only'])
        self.assertEqual(plan['sheets'][0]['name'],'Planung')
        self.assertEqual(plan['formula_count'],1)
        formula=next(cell for cell in plan['cells'] if 'formula' in cell)
        self.assertEqual(formula['cached_value_status'],'missing')
        self.assertEqual(len(plan['sha256']),64)

    def test_calendar_candidates_use_current_workbook_and_expose_schedule_collision(self):
        workbook=Workbook();details=workbook.active;details.title='Termine & Details'
        details.append(['Typ','Kurs','Teil','Termin','Beginn','Lehrende','TN','Pruefung','Thema','Code','Quelle'])
        for _ in range(4):details.append([])
        rows=[
            ['Wiederholung','Makroökonomie','Vorlesung','Mi 12–14 wöchentlich\nHS 2010',datetime(2026,10,14),'Prof. X',20,'Klausur','Thema','35-320','Q01'],
            ['Wiederholung','Makroökonomie','Übung','Mo 14–16 wöchentlich\nHS 2010',datetime(2026,10,12),'Tutor',20,'Klausur','Thema','35-320','Q02'],
            ['Wiederholung','Einführung in die Ökonometrie','Vorlesung','Mo 10–12 wöchentlich\nSFG 0150',datetime(2026,10,12),'Prof. Y',20,'Klausur','Thema','35-520','Q04'],
            ['Wiederholung','Einführung in die Ökonometrie','Übung','Mi 14–16 wöchentlich\nGW1-HS H0070',datetime(2026,10,14),'Prof. Y',20,'Klausur','Thema','35-520','Q05'],
            ['Profilseminar','Digitalisierung - Eine volkswirtschaftliche Perspektive','Seminar','Mo 10–12 wöchentlich\nWiWi1 A1020',datetime(2026,10,12),'Prof. Z',20,'Hausarbeit','Thema','35-430/440','Q08'],
            ['Projektmodul','Projektmodul Innovationsökonomik (in englischer Sprache)','Projektplenum','Do 10–12 wöchentlich\nWiWi2 F3290',datetime(2026,10,15),'Prof. A',20,'Projektarbeit','Thema','35-460','Q14'],
        ]
        for row in rows:details.append(row)
        choices=workbook.create_sheet('Kurse wählen')
        for row in range(7,7+len(rows)):
            choices.cell(row,2,rows[row-7][1]);choices.cell(row,9,'Offen')
        buffer=BytesIO();workbook.save(buffer)
        without_dimensions=BytesIO()
        with zipfile.ZipFile(BytesIO(buffer.getvalue())) as source, zipfile.ZipFile(without_dimensions,'w',zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                payload=source.read(item.filename)
                if item.filename.startswith('xl/worksheets/sheet'):
                    payload=re.sub(rb'<dimension[^>]*/>',b'',payload)
                target.writestr(item,payload)
        from openpyxl import load_workbook
        probe=load_workbook(BytesIO(without_dimensions.getvalue()),read_only=True,data_only=True)
        self.assertIsNone(probe['Termine & Details'].max_row)
        probe.close()
        with patch.object(core.backend,'cloud',return_value=without_dimensions.getvalue()):
            result=self.app.state.corpus.calendar_candidates()
        self.assertTrue(result['read_only'])
        self.assertEqual(len(result['candidates']),6)
        self.assertEqual(len(result['conflicts']),1)
        econometrics=next(item for item in result['candidates'] if item['course']=='Einführung in die Ökonometrie' and item['title'].endswith('Vorlesung'))
        self.assertEqual(econometrics['start_local'],'2026-10-12T10:00:00+02:00')
        self.assertIn('Kalenderkonflikt',econometrics['notes'])
        self.assertIn('2026-12-21',econometrics['exdates'])
        self.assertEqual(econometrics['rrule']['byday'],'MO')

    def test_embedding_work_queue_persists_content_hash_until_indexed(self):
        corpus=self.app.state.corpus
        corpus._reconcile_queue(self.path/'corpus',increment_attempts=True)
        with core.backend.db() as db:
            row=db.execute('SELECT document_id,sha256,status,attempts FROM document_queue').fetchone()
            self.assertEqual(row['sha256'],'synthetic')
            self.assertEqual(row['status'],'pending_embedding')
            self.assertEqual(row['attempts'],1)
            passage=db.execute('SELECT rowid,text FROM passages LIMIT 1').fetchone()
            key=core.embedding_api.key_for(passage['text'])
            vector=np.ones(core.embedding_api.DIMENSIONS,dtype=np.float32)
            db.execute('INSERT OR REPLACE INTO embedding_cache VALUES(?,?,?,?)',
                       (key,core.embedding_api.MODEL,core.embedding_api.DIMENSIONS,vector.tobytes()))
            db.execute('INSERT OR REPLACE INTO passage_vectors VALUES(?,?)',(passage['rowid'],key))
        corpus._reconcile_queue(self.path/'corpus')
        with core.backend.db() as db:
            state=db.execute('SELECT status FROM document_queue').fetchone()[0]
        self.assertEqual(state,'complete')

    def test_downloader_status_reads_existing_report_and_keeps_file_failures(self):
        report={'status':'partial_failure','finished':'2026-10-20T10:00:00Z','downloaded':0,
                'errors':[{'course':'pm','file':'x','error':'Stud.IP HTTP 403'}],
                'windows_scheduler_disabled':False}
        with patch.object(core.backend,'cloud',return_value=json.dumps(report).encode()):
            result=self.app.state.corpus.downloader_status()
        self.assertEqual(result['status'],'partial_failure')
        self.assertEqual(result['errors'][0]['error'],'Stud.IP HTTP 403')
        self.assertFalse(result['windows_scheduler_disabled'])

    def test_chatgpt_search_fetch_schema_and_private_citation_viewer(self):
        r=self.rpc(SERVICE,'tools/call',{'name':'search','arguments':{'query':'Instrumentvariablen'}})
        hit=r.json()['result']['structuredContent']['results'][0]
        self.assertTrue(hit['url'].startswith(URL+'/source/'))
        r=self.rpc(SERVICE,'tools/call',{'name':'fetch','arguments':{'id':hit['id']}})
        item=r.json()['result']['structuredContent']
        self.assertIn('Endogenitaetsproblem',item['text'])
        self.assertEqual(item['metadata']['page'],2)
        page=self.client.get(hit['url'])
        self.assertNotIn('Endogenitaetsproblem',page.text)
        csrf=re.search('name="csrf" value="([^"]+)"',page.text).group(1)
        r=self.client.post(hit['url'],data={'csrf':csrf,'owner_secret':OWNER},headers={'Origin':URL},follow_redirects=False)
        self.assertEqual(r.status_code,303)
        page=self.client.get(hit['url'])
        self.assertIn('Endogenitaetsproblem',page.text)

    def test_oauth_consent_pkce_code_replay_refresh_and_revoke(self):
        client=self.register();verifier,form=self.authorization(client)
        denied=self.client.post('/consent',data=form,headers={'Origin':'https://attacker.example'})
        self.assertEqual(denied.status_code,403)
        denied=self.client.post('/consent',data=dict(form,owner_secret='wrong'),headers={'Origin':URL})
        self.assertEqual(denied.status_code,403)
        ok=self.client.post('/consent',data=form,headers={'Origin':URL},follow_redirects=False)
        self.assertEqual(ok.status_code,303,ok.text)
        params=parse_qs(urlparse(ok.headers['location']).query)
        self.assertEqual(params['state'],['test-state'])
        code=params['code'][0]
        bad=self.exchange(client,code,'bad-verifier')
        self.assertIn(bad.status_code,[400,401]);self.assertEqual(bad.json()['error'],'invalid_grant')
        token=self.exchange(client,code,verifier)
        self.assertEqual(token.status_code,200,token.text)
        data=token.json()
        self.assertEqual(self.rpc(data['access_token']).status_code,200)
        self.assertIn(self.exchange(client,code,verifier).status_code,[400,401])
        refresh={'grant_type':'refresh_token','client_id':client['client_id'],'client_secret':client['client_secret'],
                 'refresh_token':data['refresh_token'],'scope':'uni:read'}
        new=self.client.post('/token',data=refresh)
        self.assertEqual(new.status_code,200,new.text)
        self.assertIn(self.client.post('/token',data=refresh).status_code,[400,401])
        revoke=self.client.post('/revoke',data={'client_id':client['client_id'],'client_secret':client['client_secret'],
            'token':new.json()['access_token'],'token_type_hint':'access_token'})
        self.assertEqual(revoke.status_code,200,revoke.text)
        self.assertEqual(self.rpc(new.json()['access_token']).status_code,401)
        raw=(self.path/'oauth.sqlite').read_bytes()
        for secret in [OWNER,SERVICE,code,data['access_token'],data['refresh_token']]:
            self.assertNotIn(secret.encode(),raw)

    def test_oauth_redirect_scope_and_resource_guards(self):
        client=self.register()
        r=self.client.get('/authorize',params={'client_id':client['client_id'],'response_type':'code',
            'redirect_uri':'https://evil.example/','code_challenge':'x'*43,'code_challenge_method':'S256'},follow_redirects=False)
        self.assertEqual(r.status_code,400)
        r=self.client.post('/register',json={'redirect_uris':['https://user:pass@client.example/callback']})
        self.assertEqual(r.status_code,400)
        r=self.client.post('/register',json={'redirect_uris':['https://client.example/'],'scope':'uni:write'})
        self.assertEqual(r.status_code,400)
        r=self.client.post('/register',json={'redirect_uris':['http://example.com/']})
        self.assertEqual(r.status_code,400)

    def test_index_update_and_removed_material_not_returned(self):
        cache=self.path/'corpus'
        doc=cache/'documents'/(self.key+'.json');data=json.loads(doc.read_text())
        data['sha256']='updated';data['chunks']=[{'page':3,'text':'Paneldaten werden aktualisiert.'}]
        doc.write_text(json.dumps(data));core.backend.ingest(cache)
        self.assertEqual(self.app.state.corpus.search('Instrumentvariablen')['sources'],[])
        self.assertTrue(self.app.state.corpus.search('Paneldaten')['sources'])
        (cache/'catalogue.json').write_text('{"files":{},"status":{"pending":0}}')
        core.backend.ingest(cache)
        self.assertEqual(self.app.state.corpus.search('Paneldaten')['sources'],[])

    def test_pending_reindex_preserves_and_labels_previous_snapshot(self):
        cache=self.path/'corpus'
        catalogue={'files':{'01_StudIP/Test.pdf':{'id':self.key,'status':'pending',
            'version':{'size':123,'extractor':2}}},'status':{'pending':1,'complete':False}}
        (cache/'catalogue.json').write_text(json.dumps(catalogue))
        core.backend.ingest(cache)
        result=self.app.state.corpus.search('Instrumentvariablen')
        self.assertEqual(result['sources'][0]['document_id'],self.key)
        self.assertEqual(result['sources'][0]['source_status'],'reindexing')
        self.assertEqual(result['reprocessing_documents'],1)
        self.assertTrue(any('vorherigen Extraktionen' in warning for warning in result['warnings']))
        self.assertEqual(self.app.state.corpus.page(self.key,2)['source_status'],'reindexing')

        with core.backend.db() as db:
            db.execute('DELETE FROM passages WHERE doc_id=?',(self.key,))
            db.execute('DELETE FROM docs WHERE id=?',(self.key,))
        core.backend.ingest(cache)
        self.assertTrue(self.app.state.corpus.search('Instrumentvariablen')['sources'])

        catalogue['files']['01_StudIP/Test.pdf']['status']='error'
        catalogue['files']['01_StudIP/Test.pdf']['error']='OCR-Pruefung erforderlich'
        (cache/'catalogue.json').write_text(json.dumps(catalogue))
        core.backend.ingest(cache)
        result=self.app.state.corpus.search('Instrumentvariablen')
        self.assertEqual(result['sources'][0]['source_status'],'refresh_error')
        self.assertEqual(result['refresh_error_documents'],1)

    def test_optional_hybrid_cache_and_api_failure_fallback(self):
        vector=np.zeros(512,dtype=np.float32);vector[0]=1
        with patch.dict(os.environ,OPENAI_API_KEY='synthetic-key'),patch.object(core.embedding_api,'request',side_effect=lambda texts,key:[vector for _ in texts]):
            with core.backend.db() as db:
                core.embedding_api.populate(db,'synthetic-key')
            result=self.app.state.corpus.search('Kausaler Zusammenhang')
            self.assertEqual(result['retrieval'],'hybrid')
            self.assertEqual(result['sources'][0]['document_id'],self.key)
        with patch.dict(os.environ,OPENAI_API_KEY='synthetic-key'),patch.object(core.embedding_api,'request',side_effect=RuntimeError('synthetic failure')):
            result=self.app.state.corpus.search('Instrumentvariablen')
            self.assertEqual(result['retrieval'],'fulltext')
            self.assertTrue(result['sources'])
            self.assertTrue(result['warnings'])

    def test_build_progress_and_safe_failure_diagnostics(self):
        corpus=self.app.state.corpus
        def failed_build(db,key,progress):
            progress({'status':'building','pending':7,'unique_chunks':10})
            self.assertEqual(corpus.last['embedding']['pending'],7)
            raise core.embedding_api.EmbeddingError('Embedding-API: HTTP 429. private-body-must-not-leak')
        with patch.object(corpus,'configure_cloud'),patch.object(core.backend,'cloud'),patch.object(core.embedding_api,'populate',side_effect=failed_build):
            corpus.synchronize()
        status=corpus.status()
        self.assertEqual(status['sync']['embedding']['http_status'],429)
        self.assertEqual(status['sync']['embedding']['pending'],7)
        self.assertEqual(status['embedding_model'],'text-embedding-3-small')
        self.assertEqual(status['embedding_dimensions'],512)
        self.assertNotIn('private-body',json.dumps(status))
        self.assertFalse(corpus.lock.locked())
        with patch.object(corpus,'configure_cloud'),patch.object(core.backend,'cloud'),patch.object(core.embedding_api,'populate',side_effect=RuntimeError('secret-value')):
            corpus.synchronize()
        self.assertIsNone(corpus.last['embedding']['http_status'])
        self.assertNotIn('secret-value',json.dumps(corpus.status()))


if __name__=='__main__':unittest.main()
