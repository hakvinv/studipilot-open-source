"""Consume the existing OneDrive corpus; never authenticate to Stud.IP."""
import base64
import hashlib
import importlib.util
import io
import json
import os
import re
import sqlite3
import subprocess
import sys
import threading
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'uni-rag'))
import app as backend
import embedding_api

spec = importlib.util.spec_from_file_location('context_helper',ROOT/'skills/uni-kontext/scripts/uni_context.py')
context = importlib.util.module_from_spec(spec)
spec.loader.exec_module(context)


def utc_stamp():
    return datetime.now(timezone.utc).isoformat()


class Corpus:
    def __init__(self,state):
        self.state=state
        self.state.mkdir(parents=True,exist_ok=True)
        backend.STATE=context.STATE=state
        backend.RCLONE=os.environ.get('RCLONE_BIN','rclone')
        self.stop=threading.Event()
        self.lock=threading.Lock()
        self.last={'state':'starting'}
        with backend.db() as db:
            embedding_api.schema(db)
            db.execute('''CREATE TABLE IF NOT EXISTS document_queue(
                document_id TEXT NOT NULL,sha256 TEXT NOT NULL,source TEXT NOT NULL,status TEXT NOT NULL,
                attempts INTEGER NOT NULL DEFAULT 0,last_error TEXT,updated_at TEXT NOT NULL,
                PRIMARY KEY(document_id,sha256))''')

    def configure_cloud(self):
        seed=os.environ.get('ONEDRIVE_CONFIG_B64')
        config=self.state/'rclone.conf'
        if seed:
            fingerprint=hashlib.sha256(seed.encode()).hexdigest()
            marker=self.state/'seed-sha256'
            if not config.exists() or not marker.exists() or marker.read_text()!=fingerprint:
                raw=base64.b64decode(seed,validate=True)
                if b'[uni]' not in raw:
                    raise ValueError('OneDrive remote uni missing')
                config.write_bytes(raw);config.chmod(0o600);marker.write_text(fingerprint)
        if not config.exists():
            raise RuntimeError('Private OneDrive configuration missing')

    def synchronize(self):
        if not self.lock.acquire(False):
            return
        try:
            self.configure_cloud()
            cache=self.state/'corpus';cache.mkdir(exist_ok=True)
            backend.cloud('copy',backend.ROOT+'/.automation/rag',str(cache),'--checksum','--transfers','8',timeout=1200)
            status=backend.ingest(cache)
            self._reconcile_queue(cache, increment_attempts=True)
            self.last={'state':'ok','finished':datetime.now(timezone.utc).isoformat(),'corpus':status}
            try:
                self.last['embedding']={'status':'starting'}
                def progress(value):
                    self.last['embedding']=dict(value)
                with backend.db() as db:
                    self.last['embedding']=embedding_api.populate(db,os.environ.get('OPENAI_API_KEY'),progress)
                self._reconcile_queue(cache, increment_attempts=False)
            except Exception as exc:
                # Only expose a numeric HTTP status, never upstream bodies or credentials.
                code=re.search(r'^Embedding-API: HTTP ([0-9]{3})\.',str(exc)) if isinstance(exc,embedding_api.EmbeddingError) else None
                self.last['embedding']={**self.last.get('embedding',{}),'status':'error',
                    'message':'Embedding-Aufbau angehalten; vorhandene Vektoren und Volltext bleiben nutzbar.',
                    'error_type':type(exc).__name__,'http_status':int(code.group(1)) if code else None,
                    'api_code':exc.api_code if isinstance(exc,embedding_api.EmbeddingError) else None,
                    'retry_after':exc.retry_after if isinstance(exc,embedding_api.EmbeddingError) else None}
                with backend.db() as db:
                    db.execute("UPDATE document_queue SET last_error=?,updated_at=? WHERE status='pending_embedding'",
                               (f'Embedding fehlgeschlagen ({type(exc).__name__})',utc_stamp()))
        except Exception:
            # Raw upstream exceptions can contain credentials or private URLs.
            self.last={'state':'error','message':'Cloudindex konnte nicht aktualisiert werden. Privaten OneDrive-Zugang und Cloudkatalog pruefen.'}
        finally:
            self.lock.release()

    def worker(self):
        while not self.stop.is_set():
            self.synchronize()
            self.stop.wait(max(300,self.last.get('embedding',{}).get('retry_after') or 0))

    def status(self):
        data=context.local_status()
        with context.connect() as db:
            count=db.execute('SELECT count(*) FROM passages JOIN passage_vectors ON passage_vectors.passage_rowid=passages.rowid JOIN embedding_cache ON embedding_cache.key=passage_vectors.content_key').fetchone()[0]
        configured=bool(os.environ.get('OPENAI_API_KEY'))
        result=context.summary(data,'hybrid' if configured and count else 'fulltext',[])
        result.update(embedding_configured=configured,embedding_available_passages=count,
                      embedding_model=embedding_api.MODEL,embedding_dimensions=embedding_api.DIMENSIONS)
        with backend.db() as db:
            freshness=db.execute('SELECT status,count(*) AS count FROM document_freshness GROUP BY status').fetchall()
            freshness_counts={row['status']:row['count'] for row in freshness}
            queue=db.execute('SELECT status,count(*) AS count FROM document_queue GROUP BY status').fetchall()
            failures=db.execute("SELECT source,last_error FROM document_queue WHERE last_error IS NOT NULL ORDER BY updated_at DESC LIMIT 20").fetchall()
        result['reprocessing_documents']=freshness_counts.get('reindexing',0)
        result['refresh_error_documents']=freshness_counts.get('refresh_error',0)
        if result['reprocessing_documents'] or result['refresh_error_documents']:
            result['warnings'].append('Ein Teil der Treffer stammt aus vorherigen Extraktionen und wird aktuell neu verarbeitet; Treffer kennzeichnen den Stand.')
        result['processing_queue']={'counts':{row['status']:row['count'] for row in queue},
                                    'errors':[{'source':row['source'],'error':row['last_error']} for row in failures]}
        result['sync']=self.last
        return result

    def _reconcile_queue(self,cache,increment_attempts=False):
        """Persist document+content-hash work state so embeddings resume after restarts."""
        catalogue_path=Path(cache)/'catalogue.json'
        try:catalogue=json.loads(catalogue_path.read_text())
        except (OSError,json.JSONDecodeError):catalogue={}
        statuses={item.get('id'):(item.get('status'),item.get('error'))
                  for item in catalogue.get('files',{}).values() if item.get('id')}
        with backend.db() as db:
            for source,item in catalogue.get('files',{}).items():
                if item.get('status')!='error' or not item.get('id') or not item.get('sha256'):
                    continue
                db.execute('''INSERT INTO document_queue VALUES(?,?,?,?,?,?,?)
                  ON CONFLICT(document_id,sha256) DO UPDATE SET source=excluded.source,status='error',
                  last_error=excluded.last_error,updated_at=excluded.updated_at''',
                  (item['id'],item['sha256'],source,'error',0,str(item.get('error') or 'Extraktion fehlgeschlagen')[:300],utc_stamp()))
            docs=db.execute('SELECT id,sha,source FROM docs').fetchall()
            for doc in docs:
                document_id,sha,source=doc['id'],doc['sha'],doc['source']
                indexed_state,index_error=statuses.get(document_id,('indexed',None))
                if indexed_state in {'error','pending'}:
                    queue_state='source_error' if indexed_state=='error' else 'pending_extraction'
                    db.execute('''INSERT INTO document_queue VALUES(?,?,?,?,?,?,?)
                      ON CONFLICT(document_id,sha256) DO UPDATE SET source=excluded.source,status=excluded.status,
                      last_error=excluded.last_error,updated_at=excluded.updated_at''',
                      (document_id,sha,source,queue_state,0,
                       str(index_error or ('Extraktion fehlgeschlagen' if indexed_state=='error' else 'Neu-Extraktion ausstehend'))[:300],utc_stamp()))
                    continue
                missing=db.execute('''SELECT count(*) FROM passages p LEFT JOIN passage_vectors v ON v.passage_rowid=p.rowid
                      LEFT JOIN embedding_cache e ON e.key=v.content_key WHERE p.doc_id=? AND e.key IS NULL''',(document_id,)).fetchone()[0]
                passage_count=db.execute('SELECT count(*) FROM passages WHERE doc_id=?',(document_id,)).fetchone()[0]
                new_state='pending_embedding' if missing else ('review_required' if indexed_state=='ocr_or_review_required' or not passage_count else 'complete')
                old=db.execute('SELECT status,attempts FROM document_queue WHERE document_id=? AND sha256=?',(document_id,sha)).fetchone()
                attempts=(old['attempts'] if old else 0)+(1 if increment_attempts and missing else 0)
                db.execute('''INSERT INTO document_queue VALUES(?,?,?,?,?,?,?)
                  ON CONFLICT(document_id,sha256) DO UPDATE SET source=excluded.source,status=excluded.status,
                  attempts=excluded.attempts,last_error=CASE WHEN excluded.status='pending_embedding' THEN document_queue.last_error ELSE NULL END,
                  updated_at=excluded.updated_at''',
                  (document_id,sha,source,new_state,attempts,None,utc_stamp()))
    def courses(self):
        with context.connect() as db:
            return [dict(r) for r in db.execute('SELECT course,count(*) AS documents FROM docs GROUP BY course ORDER BY course')]

    def search(self,question,course='',limit=6):
        if not 1<=len(question)<=1000 or not 1<=limit<=10 or len(course)>200:
            raise ValueError('question: 1..1000 characters; limit: 1..10')
        courses=context.matching([x['course'] for x in self.courses()],course)
        rows=context.lexical(question,courses)
        state=self.status()
        if state['retrieval']=='hybrid':
            try:
                with backend.db() as db:
                    semantic=[]
                    for group in courses or ['']:
                        semantic.extend(embedding_api.semantic(db,question,group,os.environ['OPENAI_API_KEY']))
                    found={r['rowid']:dict(r) for r in db.execute('SELECT passages.rowid,passages.doc_id,page,text,source,course FROM passages JOIN docs ON docs.id=passages.doc_id WHERE passages.rowid IN ('+','.join('?' for _ in semantic)+')',[x[0] for x in semantic])} if semantic else {}
                    groups=[rows,[found[x[0]] for x in semantic if x[0] in found]]
                    ranks={};items={}
                    for group in groups:
                        for rank,row in enumerate(group,1):
                            key=(row['doc_id'],row['page'],row['text']);items[key]=row
                            ranks[key]=ranks.get(key,0)+1/(60+rank)
                    rows=[items[key] for key in sorted(ranks,key=ranks.get,reverse=True)]
            except Exception as exc:
                state['retrieval']='fulltext'
                state['warnings'].append('Embedding-Abfrage fehlgeschlagen; lokale Volltexttreffer verwendet.')
                if isinstance(exc,embedding_api.EmbeddingError):
                    state['embedding_query']={'state':'error','http_status':exc.http_status,
                        'api_code':exc.api_code,'retry_after':exc.retry_after}
                else:
                    state['embedding_query']={'state':'error','error_type':type(exc).__name__}
            else:
                state['embedding_query']={'state':'ok'}
        with backend.db() as db:
            freshness=backend.freshness_for(db,[row.get('doc_id') for row in rows])
        for row in rows:
            row['source_status']=freshness.get(row.get('doc_id'),'indexed')
        if any(row.get('source_status')!='indexed' for row in rows):
            state['warnings'].append('Mindestens ein Treffer stammt aus einer vorherigen Extraktion; source_status kennzeichnet ihn.')
        hits=context.bounded(rows,limit,1600)
        for hit in hits:
            hit.pop('original_url',None);hit.pop('download_url',None)
            hit['resource_uri']=f"uni://document/{hit['document_id']}/page/{hit['page'] or 0}"
        return {**state,'matched_courses':courses,'sources':hits}

    def page(self,document_id,page=0):
        if len(document_id)!=64 or any(c not in '0123456789abcdef' for c in document_id) or not 0<=page<=100000:
            raise ValueError('Invalid document id/page')
        with context.connect() as db:
            doc=db.execute('SELECT * FROM docs WHERE id=?',(document_id,)).fetchone()
            if not doc:
                raise ValueError('Document not indexed')
            rows=db.execute('SELECT page,text FROM passages WHERE doc_id=? AND coalesce(page,0)=? ORDER BY rowid LIMIT 31',
                            (document_id,page)).fetchall()
        with backend.db() as db:
            source_status=backend.freshness_for(db,[document_id]).get(document_id,'indexed')
        if not rows:
            raise ValueError('Page not indexed')
        # Reconstruct overlap without dropping equal text from distinct pages.
        text=''
        for row in rows[:30]:
            segment=row['text'];overlap=0
            for size in range(min(len(text),len(segment),1600),0,-1):
                if text.endswith(segment[:size]):
                    overlap=size;break
            text+=segment[overlap:]
        return {'document_id':document_id,'source':doc['source'],'course':doc['course'],'sha256':doc['sha'],
                'source_status':source_status,
                'page':page,'text':text[:24000],'complete':len(rows)<=30 and len(text)<=24000,'index':self.status()}

    def documents(self,course='',offset=0,limit=30):
        if not 0<=offset<=100000 or not 1<=limit<=100:
            raise ValueError('Invalid offset/limit')
        groups=context.matching([x['course'] for x in self.courses()],course)
        where=' WHERE course IN ('+','.join('?' for _ in groups)+')' if groups else ''
        with context.connect() as db:
            count=db.execute('SELECT count(*) FROM docs'+where,groups).fetchone()[0]
            rows=db.execute('SELECT id AS document_id,source,course,sha AS sha256 FROM docs'+where+' ORDER BY course,source LIMIT ? OFFSET ?',
                            [*groups,limit,offset]).fetchall()
        with backend.db() as db:
            freshness=backend.freshness_for(db,[row['document_id'] for row in rows])
        return {'total':count,'offset':offset,'documents':[
            {**dict(row),'source_status':freshness.get(row['document_id'],'indexed')} for row in rows]}

    def downloader_status(self):
        """Read the latest report produced by the single existing Stud.IP job."""
        try:
            raw=backend.cloud('cat',backend.ROOT+'/.automation/studip/latest.json',timeout=30)
            report=json.loads(raw)
        except Exception as exc:
            return {'state':'error','message':'Der letzte Stud.IP-Laufbericht ist in OneDrive nicht lesbar.',
                    'error_type':type(exc).__name__}
        return {key:report.get(key) for key in (
            'status','mode','started','finished','courses_discovered','courses_checked','files_seen',
            'downloaded','files_uploaded','unchanged','pdf_markdown_created','pdf_review_required',
            'roundtrip_verified','errors','courses','windows_scheduler_disabled')}

    def study_plan(self,max_cells=3000):
        """Read, but never rewrite, the authoritative OneDrive workbook."""
        from openpyxl import load_workbook

        if not 100<=max_cells<=6000:
            raise ValueError('max_cells muss zwischen 100 und 6000 liegen')
        relative='Stundenplan_WiSe_2026-27.xlsx'
        data=backend.cloud('cat',backend.ROOT+'/'+relative,timeout=60)
        formula_book=load_workbook(io.BytesIO(data),read_only=True,data_only=False)
        cached_book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        sheets=[];cells=[]
        formula_count=missing_cache_count=nonempty_count=0
        def json_value(value):
            if isinstance(value,(datetime,date,time)):
                return value.isoformat()
            return value
        try:
            for formula_sheet,cached_sheet in zip(formula_book.worksheets,cached_book.worksheets):
                sheets.append({'name':formula_sheet.title,'max_row':formula_sheet.max_row,
                               'max_column':formula_sheet.max_column})
                for row in formula_sheet.iter_rows():
                    for cell in row:
                        if cell.value is None:
                            continue
                        nonempty_count+=1
                        cached=cached_sheet[cell.coordinate].value
                        item={'sheet':formula_sheet.title,'cell':cell.coordinate,'value':json_value(cell.value)}
                        if isinstance(cell.value,str) and cell.value.startswith('='):
                            formula_count+=1
                            item={'sheet':formula_sheet.title,'cell':cell.coordinate,'formula':cell.value,
                                  'cached_value':json_value(cached),
                                  'cached_value_status':'missing' if cached is None else 'stored_not_recalculated'}
                            missing_cache_count+=cached is None
                        if len(cells)<max_cells:
                            cells.append(item)
        finally:
            formula_book.close();cached_book.close()
        return {'path':relative,'sha256':hashlib.sha256(data).hexdigest(),'size_bytes':len(data),
                'read_only':True,'sheets':sheets,'cells':cells,'cells_returned':len(cells),
                'nonempty_cells':nonempty_count,'truncated':nonempty_count>max_cells,'formula_count':formula_count,
                'formula_cached_value_missing':missing_cache_count,
                'calculation_note':'openpyxl liest gespeicherte Formelwerte; Formeln werden hier nicht neu berechnet.'}

    def calendar_candidates(self):
        """Return recurring class meetings selected in chat and evidenced by the current workbook."""
        from openpyxl import load_workbook
        from zoneinfo import ZoneInfo
        local_tz=ZoneInfo('Europe/Berlin')
        relative='Stundenplan_WiSe_2026-27.xlsx'
        data=backend.cloud('cat',backend.ROOT+'/'+relative,timeout=60)
        # The shared OneDrive workbook omits worksheet dimension metadata.
        # openpyxl's read-only mode then returns max_row=None, which breaks
        # the row scans below. This small workbook is safe to load normally;
        # it is still used strictly as a read-only source and is never saved.
        workbook=load_workbook(io.BytesIO(data),read_only=False,data_only=True)
        digest=hashlib.sha256(data).hexdigest()
        semester='WiSe 2026/27'
        selected={
            ('Makroökonomie','Vorlesung'),('Makroökonomie','Übung'),
            ('Einführung in die Ökonometrie','Vorlesung'),('Einführung in die Ökonometrie','Übung'),
            ('Digitalisierung - Eine volkswirtschaftliche Perspektive','Seminar'),
            ('Projektmodul Innovationsökonomik (in englischer Sprache)','Projektplenum'),
        }
        conflicts=[]
        source=[]
        try:
            details=workbook['Termine & Details']
            choices=workbook['Kurse wählen']
            choice_state={str(choices.cell(row,2).value):str(choices.cell(row,9).value or '')
                          for row in range(7,choices.max_row+1)}
            for row in range(6,details.max_row+1):
                kind=str(details.cell(row,1).value or '').strip()
                course=str(details.cell(row,2).value or '').strip()
                part=str(details.cell(row,3).value or '').strip()
                schedule=str(details.cell(row,4).value or '').strip()
                start_value=details.cell(row,5).value
                lecturer=str(details.cell(row,6).value or '').strip()
                course_code=str(details.cell(row,10).value or '').strip()
                source_ref=str(details.cell(row,11).value or '').strip()
                key=(course,part)
                if key not in selected or not schedule or not isinstance(start_value,(datetime,date)):
                    continue
                match=re.search(r'\b(Mo|Di|Mi|Do|Fr|Sa|So)\s+(\d{1,2})[–-](\d{1,2})\b',schedule)
                room_match=re.search(r'\n([^\n]+)$',schedule)
                if not match:
                    continue
                day,start_hour,end_hour=match.groups()
                weekdays={'Mo':'MO','Di':'TU','Mi':'WE','Do':'TH','Fr':'FR','Sa':'SA','So':'SU'}
                start_date=start_value.date() if isinstance(start_value,datetime) else start_value
                if start_date.weekday()!=({'Mo':0,'Di':1,'Mi':2,'Do':3,'Fr':4,'Sa':5,'So':6}[day]):
                    continue
                start_local=datetime.combine(start_date,time(int(start_hour)),tzinfo=local_tz)
                end_local=datetime.combine(start_date,time(int(end_hour)),tzinfo=local_tz)
                title=f'{course} — {part}'
                room=room_match.group(1).strip() if room_match else ''
                choice=choice_state.get(course,'')
                source.append({
                    'source_key':f'plan:{course_code}:{source_ref}:{part.lower()}',
                    'source_hash':digest,'source_document':relative,
                    'source_location':f'Termine & Details!A{row}:K{row}; Kurse wählen!I{next((r for r in range(7,choices.max_row+1) if str(choices.cell(r,2).value or "").strip()==course), 0)}',
                    'course':course,'semester':semester,
                    'evidence':f'{schedule.replace(chr(10), "; ")}; Beginn {start_date.isoformat()}; {lecturer}; Stud.IP-Quellkennung {source_ref}.',
                    'verified':True,'title':title,'start_local':start_local.isoformat(),
                    'end_local':end_local.isoformat(),'all_day':False,'location':room,
                    'notes':('Studienplan markiert die Kurswahl als "Offen"; eine tatsaechliche Stud.IP-Anmeldung ist hier nicht nachgewiesen. '
                             if choice.lower()=='offen' else '')+'Regeltermin laut Studienplanung.',
                    'rrule':{'freq':'weekly','byday':weekdays[day],
                             'until':datetime.combine(date(2027,1,29),time(23,59),tzinfo=local_tz).isoformat()},
                    'exdates':[date(2026,12,day_num).isoformat() for day_num in range(21,32)
                               if datetime(2026,12,day_num).weekday()==({'Mo':0,'Di':1,'Mi':2,'Do':3,'Fr':4,'Sa':5,'So':6}[day])]
                             +[date(2027,1,day_num).isoformat() for day_num in range(1,4)
                               if datetime(2027,1,day_num).weekday()==({'Mo':0,'Di':1,'Mi':2,'Do':3,'Fr':4,'Sa':5,'So':6}[day])],
                    'alarm_minutes':15,
                })
            # The user requested both selected seminars and the econometrics lecture.
            # Preserve the workbook's explicit Monday 10:00–12:00 collision.
            for i,left in enumerate(source):
                left_start=datetime.fromisoformat(left['start_local'])
                left_end=datetime.fromisoformat(left['end_local'])
                for right in source[i+1:]:
                    right_start=datetime.fromisoformat(right['start_local'])
                    right_end=datetime.fromisoformat(right['end_local'])
                    if left_start.weekday()==right_start.weekday() and left_start.time()<right_end.time() and right_start.time()<left_end.time():
                        note=f'Kalenderkonflikt: Ueberschneidung mit {right["title"]} ({right_start.strftime("%H:%M")}–{right_end.strftime("%H:%M")}).'
                        left['notes']=(left['notes']+' '+note).strip()
                        right['notes']=(right['notes']+f' Kalenderkonflikt: Ueberschneidung mit {left["title"]} ({left_start.strftime("%H:%M")}–{left_end.strftime("%H:%M")}).').strip()
                        conflicts.append({'courses':[left['course'],right['course']],
                                          'day':left_start.strftime('%A'),'start':max(left_start,right_start).strftime('%H:%M'),
                                          'end':min(left_end,right_end).strftime('%H:%M')})
        finally:
            workbook.close()
        return {'source_document':relative,'source_hash':digest,'semester':semester,
                'schedule_state':'draft; course selection marked open in workbook',
                'candidates':source,'conflicts':conflicts,'read_only':True}
